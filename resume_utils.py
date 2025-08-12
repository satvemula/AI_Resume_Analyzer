import docx
import PyPDF2
import io

def extract_text_from_docx(file):
    """
    Extract text from a DOCX file.
    """
    try:
        doc = docx.Document(file)
        text = []
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text.strip())
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text.strip())
        
        return "\n".join(text)
        
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")

def extract_text_from_pdf(file):
    """
    Extract text from a PDF file.
    """
    try:
        # Create a file-like object from the uploaded file
        file_like = io.BytesIO(file.read())
        reader = PyPDF2.PdfReader(file_like)
        
        text = []
        for page_num, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text()
                if page_text.strip():
                    text.append(page_text)
            except Exception as e:
                print(f"Warning: Could not extract text from page {page_num + 1}: {e}")
                continue
        
        if not text:
            raise ValueError("No text could be extracted from the PDF")
            
        return "\n".join(text)
        
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {str(e)}")

def extract_resume_text(file, file_type):
    """
    Extract text from resume file based on file type.
    
    Args:
        file: Uploaded file object
        file_type: String indicating file type ('pdf' or 'docx')
    
    Returns:
        String containing extracted text
    """
    file_type = file_type.lower()
    
    if file_type == "docx":
        return extract_text_from_docx(file)
    elif file_type == "pdf":
        return extract_text_from_pdf(file)
    else:
        raise ValueError(f"Unsupported file type: {file_type}. Supported types are PDF and DOCX.")

def validate_extracted_text(text):
    """
    Validate that extracted text is meaningful.
    """
    if not text or not text.strip():
        return False, "No text found in the document"
    
    if len(text.strip()) < 50:
        return False, "Document appears to contain very little text"
    
    # Check for common resume indicators
    resume_indicators = [
        'experience', 'education', 'skills', 'work', 'employment',
        'university', 'college', 'degree', 'certification', 'project'
    ]
    
    text_lower = text.lower()
    found_indicators = sum(1 for indicator in resume_indicators if indicator in text_lower)
    
    if found_indicators < 2:
        return False, "Document may not be a resume (missing common resume sections)"
    
    return True, "Text extraction successful"

def clean_extracted_text(text):
    """
    Clean and normalize extracted text.
    """
    import re
    
    # Remove excessive whitespace
    text = re.sub(r'\n\s*\n', '\n\n', text)
    text = re.sub(r' +', ' ', text)
    
    # Remove special characters that might interfere with processing
    text = re.sub(r'[^\w\s\-\.\,\(\)\[\]\{\}\/\@\#\%\&\*\+\=\|\\\:\;\"\'\?\!\$]', ' ', text)
    
    return text.strip()